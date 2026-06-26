"""Camada de acesso a dados usando planilhas Excel como 'banco de dados'.

Cada arquivo .xlsx tem uma única aba com uma linha de cabeçalho (nomes de
coluna) seguida das linhas de dados. Todo acesso de leitura+escrita é
protegido por um lock de arquivo (`<arquivo>.xlsx.lock`), o que evita que dois
usuários gravem ao mesmo tempo quando o arquivo está numa pasta de rede
compartilhada.
"""
import os
from typing import Callable, Optional

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook

import config
from core.logger import get_logger

logger = get_logger(__name__)


class ArquivoEmUsoError(Exception):
    """Levantada quando não foi possível obter o lock do arquivo a tempo."""


def _lock_path(path: str) -> str:
    return path + ".lock"


def ensure_file(path: str, columns: list[str]) -> None:
    """Cria a planilha com o cabeçalho se ela ainda não existir."""
    if os.path.exists(path):
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with FileLock(_lock_path(path), timeout=config.LOCK_TIMEOUT_SECONDS):
        if os.path.exists(path):
            return
        wb = Workbook()
        ws = wb.active
        ws.append(columns)
        wb.save(path)
    logger.info("Planilha criada: %s", path)


def _with_lock(path: str, func: Callable):
    try:
        with FileLock(_lock_path(path), timeout=config.LOCK_TIMEOUT_SECONDS):
            return func()
    except Timeout as exc:
        logger.warning("Timeout aguardando lock de %s", path)
        raise ArquivoEmUsoError(
            f"O arquivo {os.path.basename(path)} está em uso por outro usuário. "
            "Tente novamente em alguns segundos."
        ) from exc


def read_all(path: str, columns: list[str]) -> list[dict]:
    """Lê todas as linhas da planilha como lista de dicts (coluna -> valor)."""
    ensure_file(path, columns)

    def _read():
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return []
        header = [str(h) if h is not None else "" for h in header]
        rows = []
        for raw_row in rows_iter:
            if raw_row is None or all(v is None for v in raw_row):
                continue
            row_dict = dict(zip(header, raw_row))
            rows.append({col: ("" if row_dict.get(col) is None else row_dict.get(col)) for col in columns})
        wb.close()
        return rows

    return _with_lock(path, _read)


def append_row(path: str, columns: list[str], row: dict) -> None:
    """Adiciona uma nova linha ao final da planilha."""
    ensure_file(path, columns)

    def _append():
        wb = load_workbook(path)
        ws = wb.active
        ws.append([row.get(col, "") for col in columns])
        wb.save(path)
        logger.info("Linha adicionada em %s: %s", os.path.basename(path), row)

    _with_lock(path, _append)


def update_row(path: str, columns: list[str], key_col: str, key_val, updates: dict) -> bool:
    """Atualiza a primeira linha cuja coluna `key_col` seja igual a `key_val`."""
    ensure_file(path, columns)
    key_idx = columns.index(key_col)

    def _update():
        wb = load_workbook(path)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            cell = row[key_idx]
            if str(cell.value) == str(key_val):
                for col_name, value in updates.items():
                    if col_name in columns:
                        col_idx = columns.index(col_name)
                        row[col_idx].value = value
                wb.save(path)
                logger.info(
                    "Linha atualizada em %s (%s=%s): %s",
                    os.path.basename(path), key_col, key_val, updates,
                )
                return True
        return False

    return _with_lock(path, _update)


def delete_row(path: str, columns: list[str], key_col: str, key_val) -> bool:
    """Remove a primeira linha cuja coluna `key_col` seja igual a `key_val`."""
    ensure_file(path, columns)
    key_idx = columns.index(key_col)

    def _delete():
        wb = load_workbook(path)
        ws = wb.active
        for row_num in range(2, ws.max_row + 1):
            if str(ws.cell(row=row_num, column=key_idx + 1).value) == str(key_val):
                ws.delete_rows(row_num)
                wb.save(path)
                logger.info(
                    "Linha removida em %s (%s=%s)",
                    os.path.basename(path), key_col, key_val,
                )
                return True
        return False

    return _with_lock(path, _delete)


def find_one(path: str, columns: list[str], key_col: str, key_val) -> Optional[dict]:
    for row in read_all(path, columns):
        if str(row.get(key_col)) == str(key_val):
            return row
    return None
