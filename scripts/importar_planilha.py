"""Importa equipamentos de uma planilha Excel existente para o sistema.

Uso:
    python scripts/importar_planilha.py "C:\\caminho\\planilha_antiga.xlsx" --listar-colunas
        -> apenas mostra os cabeçalhos encontrados na planilha (nenhum dado é gravado).

    python scripts/importar_planilha.py "C:\\caminho\\planilha_antiga.xlsx"
        -> importa os equipamentos para data/equipamentos.xlsx.

    python scripts/importar_planilha.py "C:\\caminho\\planilha_antiga.xlsx" --sobrescrever
        -> se um código já existir no sistema, atualiza os dados em vez de pular.

Antes de importar de verdade, rode com --listar-colunas e confira se o
MAPEAMENTO abaixo bate com os cabeçalhos reais da sua planilha. Ajuste o
dicionário MAPEAMENTO conforme necessário (a chave é o texto do cabeçalho na
planilha antiga, o valor é o nome do campo no sistema).
"""
import argparse
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

import config
from core import excel_db
from core.models import Equipamento, StatusEquipamento

# ─── Mapeamento: cabeçalho da planilha antiga -> campo do sistema ────────────
# Ajuste os textos à esquerda para bater com os cabeçalhos reais da sua
# planilha (não precisa ser exato: acentos/maiúsculas são ignorados na
# comparação, mas o texto em si precisa corresponder).
MAPEAMENTO = {
    "codigo": "codigo",
    "código": "codigo",
    "instrumento": "codigo",
    "grupo": "grupo",
    "certificado": "certificado",
    "descricao": "descricao",
    "descrição": "descricao",
    "equipamento": "descricao",
    "numero de serie": "numero_serie",
    "número de série": "numero_serie",
    "n serie": "numero_serie",
    "nº série": "numero_serie",
    "n° serie": "numero_serie",
    "fabricante": "fabricante",
    "fornecedor": "fornecedor",
    "status": "status",
    "situacao": "status",
    "situação": "status",
    "descricao status": "descricao_status",
    "descrição do status": "descricao_status",
    "periodicidade": "periodicidade_calibracao_meses",
    "periodicidade calibracao meses": "periodicidade_calibracao_meses",
    "periodicidade de calibracao (meses)": "periodicidade_calibracao_meses",
    "ultima calibracao": "data_ultima_calibracao",
    "última calibração": "data_ultima_calibracao",
    "data ultima calibracao": "data_ultima_calibracao",
    "proxima calibracao": "data_proxima_calibracao",
    "próxima calibração": "data_proxima_calibracao",
    "data proxima calibracao": "data_proxima_calibracao",
    "vencimento": "data_proxima_calibracao",
    "localizacao": "localizacao",
    "localização": "localizacao",
    "instalacao": "instalacao",
    "instalação": "instalacao",
    "condicionamento": "condicionamento",
    "modelo": "modelo",
    "local de calibracao": "local_calibracao",
    "local de calibração": "local_calibracao",
}

# Normaliza texto de status da planilha antiga para os status aceitos pelo sistema.
MAPEAMENTO_STATUS = {
    "disponivel": StatusEquipamento.DISPONIVEL,
    "em uso": StatusEquipamento.EM_USO,
    "emprestado": StatusEquipamento.EMPRESTADO,
    "calibracao vencida": StatusEquipamento.CALIBRACAO_VENCIDA,
    "vencida": StatusEquipamento.CALIBRACAO_VENCIDA,
    "aguardando calibracao": StatusEquipamento.AGUARDANDO_CALIBRACAO,
    "fora de uso": StatusEquipamento.FORA_DE_USO,
    "descalibrado": StatusEquipamento.DESCALIBRADO,
}

_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y")


def _normalizar(texto: str) -> str:
    texto = str(texto or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return " ".join(texto.split())


# Versão do MAPEAMENTO com as chaves normalizadas, usada nas buscas (evita que
# acentos/maiúsculas/símbolos façam a comparação falhar).
_MAPEAMENTO_NORMALIZADO = {_normalizar(k): v for k, v in MAPEAMENTO.items()}


def _formatar_data(valor) -> str:
    if valor in (None, ""):
        return ""
    if isinstance(valor, datetime):
        return valor.date().strftime("%Y-%m-%d")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    texto = str(valor).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(texto, fmt).date().strftime("%Y-%m-%d")
        except ValueError:
            continue
    print(f"  [aviso] data não reconhecida, mantendo como texto: {valor!r}")
    return texto


def _formatar_status(valor) -> str:
    if valor in (None, ""):
        return StatusEquipamento.DISPONIVEL
    chave = _normalizar(valor)
    return MAPEAMENTO_STATUS.get(chave, str(valor).strip())


def listar_colunas(caminho: str, aba: str | None) -> None:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    print(f"Abas encontradas: {wb.sheetnames}")
    ws = wb[aba] if aba else wb.active
    print(f"Usando aba: {ws.title}")
    header = next(ws.iter_rows(values_only=True), None)
    if not header:
        print("Planilha vazia.")
        return
    print("\nCabeçalhos encontrados (copie/ajuste no MAPEAMENTO do script):")
    for h in header:
        if h is None:
            continue
        h = str(h).strip()
        campo = _MAPEAMENTO_NORMALIZADO.get(_normalizar(h), "??? (sem mapeamento -> será ignorado)")
        print(f"  {h!r:45s} -> {campo}")
    wb.close()


def importar(caminho: str, aba: str | None, sobrescrever: bool) -> None:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba] if aba else wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header_bruto = next(rows_iter, None)
    if not header_bruto:
        print("Planilha vazia, nada para importar.")
        return

    indice_campo = {}
    for idx, h in enumerate(header_bruto):
        if h is None:
            continue
        campo = _MAPEAMENTO_NORMALIZADO.get(_normalizar(h))
        if campo:
            indice_campo[campo] = idx

    if "codigo" not in indice_campo:
        print(
            "[ERRO] Não encontrei uma coluna de 'código' na planilha. "
            "Rode com --listar-colunas e ajuste o MAPEAMENTO no início do script."
        )
        return

    existentes = {
        row["codigo"]: row
        for row in excel_db.read_all(config.EQUIPAMENTOS_FILE, Equipamento.colunas())
    }

    importados = 0
    atualizados = 0
    pulados = 0
    for num_linha, raw_row in enumerate(rows_iter, start=2):
        if raw_row is None or all(v is None for v in raw_row):
            continue

        dados = {"status": StatusEquipamento.DISPONIVEL}
        for campo, idx in indice_campo.items():
            valor = raw_row[idx] if idx < len(raw_row) else None
            if campo in ("data_ultima_calibracao", "data_proxima_calibracao"):
                dados[campo] = _formatar_data(valor)
            elif campo == "status":
                dados[campo] = _formatar_status(valor)
            else:
                dados[campo] = "" if valor is None else str(valor).strip()

        codigo = dados.get("codigo", "").strip()
        if not codigo:
            print(f"  [linha {num_linha}] sem código, pulando.")
            pulados += 1
            continue

        equipamento = Equipamento(codigo=codigo)
        linha_completa = {**{c: getattr(equipamento, c) for c in Equipamento.colunas()}, **dados}

        if codigo in existentes:
            if sobrescrever:
                excel_db.update_row(
                    config.EQUIPAMENTOS_FILE, Equipamento.colunas(), "codigo", codigo, linha_completa
                )
                atualizados += 1
                print(f"  [linha {num_linha}] {codigo}: atualizado.")
            else:
                print(f"  [linha {num_linha}] {codigo}: já existe, pulado (use --sobrescrever para atualizar).")
                pulados += 1
            continue

        excel_db.append_row(config.EQUIPAMENTOS_FILE, Equipamento.colunas(), linha_completa)
        existentes[codigo] = linha_completa
        importados += 1
        print(f"  [linha {num_linha}] {codigo}: importado.")

    wb.close()
    print("\n─── Resumo ───────────────────────────────")
    print(f"Importados novos : {importados}")
    print(f"Atualizados      : {atualizados}")
    print(f"Pulados          : {pulados}")
    print(f"Arquivo destino  : {config.EQUIPAMENTOS_FILE}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("planilha", help="Caminho da planilha Excel antiga (.xlsx)")
    parser.add_argument("--aba", default=None, help="Nome da aba a usar (padrão: primeira aba ativa)")
    parser.add_argument("--listar-colunas", action="store_true", help="Apenas mostra os cabeçalhos encontrados, sem importar")
    parser.add_argument("--sobrescrever", action="store_true", help="Atualiza equipamentos já cadastrados em vez de pular")
    args = parser.parse_args()

    if not Path(args.planilha).exists():
        print(f"[ERRO] Arquivo não encontrado: {args.planilha}")
        sys.exit(1)

    if args.listar_colunas:
        listar_colunas(args.planilha, args.aba)
    else:
        importar(args.planilha, args.aba, args.sobrescrever)


if __name__ == "__main__":
    main()
